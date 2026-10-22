#!/usr/bin/env python
"""End to end check of the dataflow layer (import preview -> commit -> report).

The script builds a throw away SQLite database in a temporary data directory,
seeds a project with one website / task / verified backlink per scenario and
runs the whole import / export flow against it.

Run it from ``apps/api``::

    ../../.venv/bin/python scripts/check_dataflow.py
"""

from __future__ import annotations

import os
import sys
import tempfile
from datetime import timedelta
from pathlib import Path

# The managed data directory must be redirected *before* the app is imported:
# ``app.core.config`` creates its settings singleton at import time.
WORKSPACE = Path(tempfile.mkdtemp(prefix="autobl-dataflow-"))
os.environ["AUTOBL_DATA_DIR"] = str(WORKSPACE)
os.environ.setdefault("AUTOBL_ENVIRONMENT", "test")

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from app.db.base import Base, utcnow
from app.db.models import (
    BacklinkRecord,
    ExportJob,
    ImportJob,
    ImportRow,
    Project,
    Task,
    TaskAttempt,
    TaskEvent,
    User,
    VerificationRun,
    Website,
)
from app.services.dataflow.exporter import generate_report
from app.services.dataflow.importer import commit_import, preview_import
from app.services.opportunities import create_website, link_to_project
from app.services.seed import seed_registry

HEADERS = [
    "Project",
    "Website",
    "Target URL",
    "Backlink Website",
    "Submission URL",
    "Category",
    "Anchor Text",
    "Description",
    "Email",
    "Status",
    "Backlink Published URL",
    "Verified",
    "Last Attempt",
    "Notes",
    "Error",
]

PROJECT_NAME = "Acme Growth"
TARGET_URL = "https://acme.example.com/landing"


def banner(title: str) -> None:
    print()
    print("=" * 78)
    print(title)
    print("=" * 78)


# --------------------------------------------------------------------------
# Seeding
# --------------------------------------------------------------------------


def seed_base(db: Session) -> Project:
    """Create the user, the category registry and the project."""
    user = User(email=f"owner-{PROJECT_NAME.lower().replace(' ', '-')}@example.com",
                password_hash="x", name="Owner")
    db.add(user)
    db.commit()
    seed_registry(db)
    project = Project(
        name=PROJECT_NAME,
        slug="acme-growth",
        description="Dataflow check project",
        owner_id=user.id,
        target_website="https://acme.example.com",
    )
    db.add(project)
    db.commit()
    db.refresh(project)
    return project


def seed_scenarios(db: Session, project: Project, suffix: str) -> dict[str, dict]:
    """Create one website / opportunity / task per import scenario."""
    scenarios: dict[str, dict] = {}

    def website(url: str, name: str) -> Website:
        row, _created = create_website(
            db,
            url=url,
            name=name,
            category="BUSINESS_DIRECTORY",
            discovery_source="manual",
        )
        return row

    # -- UPDATE: an opportunity without a backlink record -----------------
    update_site = website(f"https://update-{suffix}.example.net/listing", "Update Site")
    update_opp, _ = link_to_project(db, project, update_site, target_landing_page=TARGET_URL)
    update_task = Task(
        project_id=project.id,
        opportunity_id=update_opp.id,
        website_id=update_site.id,
        type="SUBMIT",
        status="QUEUED",
        idempotency_key=f"seed:update:{suffix}",
        payload={},
        attempts=0,
        max_attempts=3,
    )
    db.add(update_task)
    db.commit()
    scenarios["update"] = {"website": update_site, "opportunity": update_opp, "task": update_task}

    # -- SKIP: a verified and still scheduled backlink --------------------
    verified_site = website(f"https://verified-{suffix}.example.net/page", "Verified Site")
    verified_opp, _ = link_to_project(
        db, project, verified_site, target_landing_page=TARGET_URL
    )
    verified_task = Task(
        project_id=project.id,
        opportunity_id=verified_opp.id,
        website_id=verified_site.id,
        type="SUBMIT",
        status="VERIFIED_LIVE",
        idempotency_key=f"seed:verified:{suffix}",
        payload={},
        attempts=1,
        max_attempts=3,
    )
    db.add(verified_task)
    db.commit()
    record = BacklinkRecord(
        project_id=project.id,
        opportunity_id=verified_opp.id,
        task_id=verified_task.id,
        target_website=verified_site.opportunity_url,
        expected_link_url=verified_site.opportunity_url,
        published_page_url=verified_site.opportunity_url,
        expected_destination_url=TARGET_URL,
        observed_destination_url=TARGET_URL,
        observed_anchor_text="Acme Growth",
        link_present=True,
        destination_matched=True,
        is_live=True,
        link_attribute="dofollow",
        http_status=200,
        verification_status="VERIFIED_LIVE",
        verified_at=utcnow(),
        last_success_at=utcnow(),
        next_verification_at=utcnow() + timedelta(days=7),
        evidence={"source": "seed"},
    )
    db.add(record)
    db.commit()
    # A verification run (the history of the link) and an activity event.
    db.add(
        VerificationRun(
            backlink_id=record.id,
            status="VERIFIED_LIVE",
            link_present=True,
            destination_matched=True,
            http_status=200,
            observed_anchor_text="Acme Growth",
            observed_destination_url=TARGET_URL,
            link_attribute="dofollow",
            evidence={"source": "seed"},
        )
    )
    db.add(
        TaskEvent(
            task_id=verified_task.id,
            project_id=project.id,
            website_id=verified_site.id,
            event_type="task.completed",
            message="Backlink verified live",
            level="SUCCESS",
            execution_environment="CLOUD",
            metadata={"source": "seed"},
        )
    )
    db.commit()
    # A second, still queued task: the import must mark it as SKIPPED (the
    # verified task above must not be touched).
    queued_task = Task(
        project_id=project.id,
        opportunity_id=verified_opp.id,
        website_id=verified_site.id,
        type="VERIFY",
        status="QUEUED",
        idempotency_key=f"seed:verified-queued:{suffix}",
        payload={},
        attempts=0,
        max_attempts=3,
    )
    db.add(queued_task)
    db.commit()
    scenarios["skip"] = {
        "website": verified_site,
        "opportunity": verified_opp,
        "task": verified_task,
        "queued_task": queued_task,
        "record": record,
    }

    # -- RETRY: a failed task with attempts left --------------------------
    failed_site = website(f"https://failed-{suffix}.example.net/listing", "Failed Site")
    failed_opp, _ = link_to_project(db, project, failed_site, target_landing_page=TARGET_URL)
    failed_task = Task(
        project_id=project.id,
        opportunity_id=failed_opp.id,
        website_id=failed_site.id,
        type="SUBMIT",
        status="FAILED",
        idempotency_key=f"seed:failed:{suffix}",
        payload={},
        attempts=1,
        max_attempts=3,
        last_error="Connection reset by peer",
    )
    db.add(failed_task)
    db.commit()
    db.add(
        TaskAttempt(
            task_id=failed_task.id,
            attempt_number=1,
            outcome="FAILED",
            error="Connection reset by peer",
            evidence={"http_status": 500},
        )
    )
    db.commit()
    scenarios["retry"] = {"website": failed_site, "opportunity": failed_opp, "task": failed_task}

    # -- VERIFY: a published but unverified backlink ----------------------
    published_site = website(f"https://published-{suffix}.example.net/page", "Published Site")
    published_opp, _ = link_to_project(
        db, project, published_site, target_landing_page=TARGET_URL
    )
    published_task = Task(
        project_id=project.id,
        opportunity_id=published_opp.id,
        website_id=published_site.id,
        type="SUBMIT",
        status="SUBMITTED",
        idempotency_key=f"seed:published:{suffix}",
        payload={},
        attempts=1,
        max_attempts=3,
    )
    db.add(published_task)
    db.commit()
    published_record = BacklinkRecord(
        project_id=project.id,
        opportunity_id=published_opp.id,
        task_id=published_task.id,
        target_website=published_site.opportunity_url,
        expected_link_url=published_site.opportunity_url,
        published_page_url=published_site.opportunity_url,
        expected_destination_url=TARGET_URL,
        verification_status="PUBLISHED_UNVERIFIED",
    )
    db.add(published_record)
    db.commit()
    scenarios["verify"] = {
        "website": published_site,
        "opportunity": published_opp,
        "task": published_task,
        "record": published_record,
    }
    return scenarios


def scenario_rows(suffix: str, scenarios: dict[str, dict]) -> list[list[str]]:
    """Build one row per import action."""
    return [
        # CREATE - unknown website, blank status
        [
            PROJECT_NAME,
            "Brand New Directory",
            TARGET_URL,
            f"https://new-{suffix}.example.net/submit",
            f"https://new-{suffix}.example.net/submit?form=1",
            "BUSINESS_DIRECTORY",
            "Acme Growth",
            "Acme Growth helps B2B teams grow.",
            "hello@acme.example.com",
            "",
            "",
            "",
            "",
            "imported by hand",
            "",
        ],
        # UPDATE - known opportunity, pending status
        [
            PROJECT_NAME,
            "Update Site",
            TARGET_URL,
            scenarios["update"]["website"].opportunity_url,
            "",
            "BUSINESS_DIRECTORY",
            "Acme Growth",
            "Updated description",
            "hello@acme.example.com",
            "PENDING",
            "",
            "",
            "",
            "refresh me",
            "",
        ],
        # SKIP - verified live and still scheduled
        [
            PROJECT_NAME,
            "Verified Site",
            TARGET_URL,
            scenarios["skip"]["website"].opportunity_url,
            "",
            "BUSINESS_DIRECTORY",
            "Acme Growth",
            "",
            "hello@acme.example.com",
            "VERIFIED_LIVE",
            scenarios["skip"]["website"].opportunity_url,
            "yes",
            "2026-10-01 09:00:00",
            "",
            "",
        ],
        # RETRY - failed task with attempts left
        [
            PROJECT_NAME,
            "Failed Site",
            TARGET_URL,
            scenarios["retry"]["website"].opportunity_url,
            "",
            "BUSINESS_DIRECTORY",
            "Acme Growth",
            "",
            "hello@acme.example.com",
            "FAILED",
            "",
            "",
            "2026-10-02 09:00:00",
            "",
            "Connection reset by peer",
        ],
        # VERIFY - published but unverified
        [
            PROJECT_NAME,
            "Published Site",
            TARGET_URL,
            scenarios["verify"]["website"].opportunity_url,
            "",
            "BUSINESS_DIRECTORY",
            "Acme Growth",
            "",
            "hello@acme.example.com",
            "PUBLISHED_UNVERIFIED",
            scenarios["verify"]["website"].opportunity_url,
            "no",
            "2026-10-03 09:00:00",
            "",
            "",
        ],
        # CONFLICT - the sheet claims VERIFIED_LIVE but nothing is stored
        [
            PROJECT_NAME,
            "Ghost Site",
            TARGET_URL,
            f"https://ghost-{suffix}.example.net/page",
            "",
            "BUSINESS_DIRECTORY",
            "Acme Growth",
            "",
            "hello@acme.example.com",
            "VERIFIED_LIVE",
            f"https://ghost-{suffix}.example.net/page",
            "yes",
            "",
            "",
            "",
        ],
        # INVALID - missing the required backlink website
        [
            PROJECT_NAME,
            "Broken Row",
            TARGET_URL,
            "",
            "",
            "BUSINESS_DIRECTORY",
            "",
            "",
            "hello@acme.example.com",
            "PENDING",
            "",
            "",
            "",
            "",
            "",
        ],
        # INVALID - malformed URL and a broken email address
        [
            PROJECT_NAME,
            "Broken Row",
            TARGET_URL,
            "not a valid url",
            "",
            "BUSINESS_DIRECTORY",
            "",
            "",
            "not-an-email",
            "PENDING",
            "",
            "",
            "",
            "",
            "",
        ],
    ]


# --------------------------------------------------------------------------
# File builders
# --------------------------------------------------------------------------


def build_xlsx(path: Path, rows: list[list[str]]) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Backlinks"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    return path


def build_csv(path: Path, rows: list[list[str]]) -> Path:
    import csv

    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(HEADERS)
        for row in rows:
            writer.writerow(row)
    return path


# --------------------------------------------------------------------------
# Checks
# --------------------------------------------------------------------------


def run_import(
    db: Session,
    project: Project,
    source: Path,
    filename: str,
    expected_actions: list[str],
    apply_actions: list[str] | None = None,
    dry_run: bool = False,
    expected_counters: dict[str, int] | None = None,
) -> tuple[ImportJob, dict]:
    """Preview, assert the actions, commit and return the job."""
    preview = preview_import(db, source, filename, project.id, project.owner_id or "")

    banner(f"PREVIEW - {filename}")
    print(f"import_job_id : {preview['import_job_id']}")
    print(f"sheets        : {preview['sheets']}")
    print("mapping       :")
    for entry in preview["detected_mapping"]:
        print(
            f"    {entry['source_column']:<24} -> {entry['target_field'] or '(unmapped)':<18}"
            f" confidence={entry['confidence']:<5} required={entry['required']}"
        )
    print("rows          :")
    for row in preview["rows"]:
        print(
            f"    row {row['row_number']:>3}  {row['action']:<8} {row['reason'][:80]}"
        )
    print(f"counts        : {preview['counts']}")

    actions = [row["action"] for row in preview["rows"]]
    if actions != expected_actions:
        raise AssertionError(
            f"unexpected actions for {filename}: {actions} != {expected_actions}"
        )
    if preview["counts"]["total"] != len(expected_actions):
        raise AssertionError("unexpected total count")
    check_contract(preview)

    job = db.get(ImportJob, preview["import_job_id"])
    banner(f"COMMIT - {filename} (dry_run={dry_run}, apply={apply_actions or 'ALL'})")
    commit_import(db, job, apply_actions, dry_run=dry_run)
    db.refresh(job)
    print(
        "counters      : "
        f"created={job.created_count} updated={job.updated_count} "
        f"skipped={job.skipped_count} retried={job.retried_count} "
        f"verify={job.verify_count} conflict={job.conflict_count} "
        f"invalid={job.invalid_count} status={job.status}"
    )
    print(f"error report  : {job.error_report_path}")
    if expected_counters:
        actual = {
            "created": job.created_count,
            "updated": job.updated_count,
            "skipped": job.skipped_count,
            "retried": job.retried_count,
            "verify": job.verify_count,
            "conflict": job.conflict_count,
            "invalid": job.invalid_count,
        }
        if actual != expected_counters:
            raise AssertionError(f"unexpected counters: {actual} != {expected_counters}")
    return job, preview


def check_state(db: Session, suffix: str, applied: bool) -> None:
    """Assert that the commit did (or did not) touch the database."""
    banner(f"STATE - {suffix} (applied={applied})")
    task = db.execute(
        select(Task).where(Task.idempotency_key == f"seed:failed:{suffix}")
    ).scalars().first()
    attempts = db.execute(
        select(TaskAttempt).where(TaskAttempt.task_id == task.id)
    ).scalars().all()
    print(f"failed task   : status={task.status} attempts={task.attempts}/{task.max_attempts}")
    print(f"attempts kept : {len(attempts)}")
    if applied:
        assert task.status == "QUEUED", f"expected a requeued task, got {task.status}"
        assert task.attempts == 1, "the retry must not reset the attempt counter"
        assert len(attempts) == 1, "the attempt history must not be deleted"
    else:
        assert task.status == "FAILED", f"a dry run must not change the task, got {task.status}"

    verified = db.execute(
        select(BacklinkRecord).where(
            BacklinkRecord.expected_link_url
            == f"https://verified-{suffix}.example.net/page"
        )
    ).scalars().first()
    print(
        f"verified link : status={verified.verification_status} "
        f"next={verified.next_verification_at}"
    )
    assert verified.verification_status == "VERIFIED_LIVE"
    assert verified.next_verification_at is not None

    published = db.execute(
        select(BacklinkRecord).where(
            BacklinkRecord.expected_link_url == f"https://published-{suffix}.example.net/page"
        )
    ).scalars().first()
    reimport_tasks = db.execute(
        select(Task).where(Task.idempotency_key == f"VERIFY:{published.id}:reimport")
    ).scalars().all()
    print(f"reimport VERIFY tasks : {len(reimport_tasks)} (record {published.id})")
    if applied:
        assert len(reimport_tasks) == 1, "a VERIFY task must have been enqueued"
    else:
        assert not reimport_tasks, "a dry run must not enqueue a VERIFY task"

    skipped = db.execute(
        select(Task).where(Task.idempotency_key == f"seed:verified:{suffix}")
    ).scalars().first()
    print(f"verified task : status={skipped.status}")
    if applied:
        assert (
            skipped.status == "VERIFIED_LIVE"
        ), f"a verified task must not be downgraded, got {skipped.status}"
        queued = db.execute(
            select(Task).where(Task.idempotency_key == f"seed:verified-queued:{suffix}")
        ).scalars().first()
        print(f"queued task   : status={queued.status}")
        assert queued.status == "SKIPPED", f"expected SKIPPED, got {queued.status}"


def check_error_report(job: ImportJob) -> None:
    """The error report must list every row that was not applied."""
    banner("ERROR REPORT")
    path = Path(job.error_report_path or "")
    print(f"path  : {path}")
    if not path.exists():
        raise AssertionError("the error report was not written")
    lines = path.read_text(encoding="utf-8").splitlines()
    print(f"lines : {len(lines)}")
    for line in lines[:8]:
        print(f"    {line[:140]}")


def check_unsupported(db: Session, project: Project) -> None:
    """A legacy ``.xls`` file must fail with a clear message (never crash)."""
    banner("UNSUPPORTED FILE")
    legacy = WORKSPACE / "legacy.xls"
    legacy.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1legacy")
    try:
        preview_import(db, legacy, "legacy.xls", project.id, project.owner_id or "")
    except ValueError as exc:
        print(f"raised        : ValueError({exc})")
    else:
        raise AssertionError("a legacy .xls file must raise a ValueError")
    failed = db.execute(select(ImportJob).where(ImportJob.status == "FAILED")).scalars().all()
    print(f"failed jobs   : {len(failed)}")
    if not failed:
        raise AssertionError("the failure must be recorded as a failed import job")
    print(f"job error     : {failed[-1].error}")


def check_multi_sheet(db: Session, project: Project) -> None:
    """Every sheet of a workbook must be imported."""
    banner("MULTI SHEET WORKBOOK")
    workbook = Workbook()
    first = workbook.active
    first.title = "Sheet One"
    first.append(HEADERS)
    first.append(
        [
            PROJECT_NAME,
            "Multi One",
            TARGET_URL,
            f"https://multi-one-{project.id[:6]}.example.net/page",
            "",
            "BUSINESS_DIRECTORY",
            "Acme Growth",
            "",
            "hello@acme.example.com",
            "",
            "",
            "",
            "",
            "",
            "",
        ]
    )
    second = workbook.create_sheet("Sheet Two")
    second.append(HEADERS)
    second.append(
        [
            PROJECT_NAME,
            "Multi Two",
            TARGET_URL,
            f"https://multi-two-{project.id[:6]}.example.net/page",
            "",
            "BUSINESS_DIRECTORY",
            "Acme Growth",
            "",
            "hello@acme.example.com",
            "",
            "",
            "",
            "",
            "",
            "",
        ]
    )
    path = WORKSPACE / "multi.xlsx"
    workbook.save(path)

    preview = preview_import(db, path, "multi.xlsx", project.id, project.owner_id or "")
    print(f"sheets        : {preview['sheets']}")
    print(f"rows          : {preview['counts']['total']}")
    if preview["sheets"] != ["Sheet One", "Sheet Two"]:
        raise AssertionError(f"unexpected sheets: {preview['sheets']}")
    if preview["counts"]["total"] != 2:
        raise AssertionError("both sheets must be imported")
    if preview["counts"]["create"] != 2:
        raise AssertionError("both rows must be planned as CREATE")


def check_contract(preview: dict) -> None:
    """The preview payload must match the shared DTO contract."""
    from app.schemas import ImportPreviewOut

    ImportPreviewOut(**preview)


def check_report(db: Session, project: Project) -> None:
    """Generate a workspace report and print what it contains."""
    banner("REPORT - workspace")
    job = ExportJob(
        project_id=None,
        user_id=project.owner_id,
        scope="workspace",
        status="RUNNING",
        format="xlsx",
    )
    db.add(job)
    db.commit()
    db.refresh(job)
    generate_report(db, job)
    db.refresh(job)

    path = Path(job.file_path or "")
    print(f"file name     : {job.file_name}")
    print(f"file path     : {job.file_path}")
    print(f"size          : {path.stat().st_size if path.exists() else 0} bytes")
    print(f"rows          : {job.row_count}")
    print(f"status        : {job.status}")
    workbook = load_workbook(path)
    print(f"sheets        : {workbook.sheetnames}")
    expected = [
        "Executive Summary",
        "All Opportunities",
        "Per website",
        "Verified Backlinks",
        "Pending Tasks",
        "Failed Tasks",
        "Skipped Tasks",
        "Verification History",
        "Activity Log",
        "Import Errors",
    ]
    if workbook.sheetnames != expected:
        raise AssertionError(f"unexpected sheet names: {workbook.sheetnames}")
    overview = workbook["All Opportunities"]
    print(f"header        : {[cell.value for cell in overview[1]][:6]} ...")
    print(f"data rows     : {overview.max_row - 1}")
    summary = workbook["Executive Summary"]
    print("totals row    : ", [cell.value for cell in summary[summary.max_row]])
    for name in workbook.sheetnames:
        sheet = workbook[name]
        print(f"    {name:<22} rows={sheet.max_row - 1:>3} cols={sheet.max_column}")
    history = workbook["Verification History"]
    if history.max_row > 1:
        print(
            "history sample: ",
            [cell.value for cell in history[2]][:1]
            + [history.cell(row=2, column=12).value]
            + [history.cell(row=2, column=20).value],
        )
    workbook.close()

    # CSV export (flat version of the "All Opportunities" sheet)
    csv_job = ExportJob(
        project_id=project.id,
        user_id=project.owner_id,
        scope=f"project:{project.id}",
        status="RUNNING",
        format="csv",
    )
    db.add(csv_job)
    db.commit()
    db.refresh(csv_job)
    generate_report(db, csv_job)
    db.refresh(csv_job)
    csv_path = Path(csv_job.file_path or "")
    print(f"csv file      : {csv_job.file_name} ({csv_path.stat().st_size} bytes)")


def check_history(db: Session) -> None:
    """The retry must not have destroyed the attempt history."""
    banner("HISTORY")
    attempts = db.execute(select(TaskAttempt)).scalars().all()
    print(f"task attempts kept : {len(attempts)}")
    rows = db.execute(select(ImportRow)).scalars().all()
    print(f"import rows kept   : {len(rows)}")
    for row in rows:
        print(f"    row {row.row_number:>3} {row.action:<8} {str(row.reason)[:70]}")


def main() -> int:
    banner("SETUP")
    print(f"temporary data dir : {WORKSPACE}")
    engine = create_engine(f"sqlite:///{WORKSPACE / 'check.db'}")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        project = seed_base(db)
        print(f"project            : {project.name} ({project.id})")

        scenarios_x = seed_scenarios(db, project, "xlsx")
        rows_x = scenario_rows("xlsx", scenarios_x)
        xlsx_path = build_xlsx(WORKSPACE / "backlinks.xlsx", rows_x)
        scenarios_c = seed_scenarios(db, project, "csv")
        rows_c = scenario_rows("csv", scenarios_c)
        csv_path = build_csv(WORKSPACE / "backlinks.csv", rows_c)
        scenarios_p = seed_scenarios(db, project, "partial")
        rows_p = scenario_rows("partial", scenarios_p)
        partial_path = build_xlsx(WORKSPACE / "partial.xlsx", rows_p)

        expected = [
            "CREATE",
            "UPDATE",
            "SKIP",
            "RETRY",
            "VERIFY",
            "CONFLICT",
            "INVALID",
            "INVALID",
        ]

        # 1. dry run: nothing may be written to the domain tables
        job, _preview = run_import(
            db,
            project,
            xlsx_path,
            "backlinks.xlsx",
            expected,
            dry_run=True,
            expected_counters={
                "created": 1,
                "updated": 1,
                "skipped": 1,
                "retried": 1,
                "verify": 1,
                "conflict": 1,
                "invalid": 2,
            },
        )
        if job.status != "PREVIEWED":
            raise AssertionError(f"a dry run must not commit the job, got {job.status}")
        check_state(db, "xlsx", applied=False)

        # 2. the real commit applies every planned action
        job, _preview = run_import(
            db,
            project,
            xlsx_path,
            "backlinks.xlsx",
            expected,
            expected_counters={
                "created": 1,
                "updated": 1,
                "skipped": 1,
                "retried": 1,
                "verify": 1,
                "conflict": 1,
                "invalid": 2,
            },
        )
        check_state(db, "xlsx", applied=True)
        check_error_report(job)

        # 3. partial commit: only the requested actions are applied
        job, _preview = run_import(
            db,
            project,
            partial_path,
            "partial.xlsx",
            expected,
            apply_actions=["CREATE", "UPDATE"],
            expected_counters={
                "created": 1,
                "updated": 1,
                "skipped": 0,
                "retried": 0,
                "verify": 0,
                "conflict": 1,
                "invalid": 2,
            },
        )
        check_state(db, "partial", applied=False)
        check_error_report(job)

        # 4. the CSV import (same scenarios, different URLs)
        job, _preview = run_import(
            db,
            project,
            csv_path,
            "backlinks.csv",
            expected,
            expected_counters={
                "created": 1,
                "updated": 1,
                "skipped": 1,
                "retried": 1,
                "verify": 1,
                "conflict": 1,
                "invalid": 2,
            },
        )
        check_state(db, "csv", applied=True)

        check_report(db, project)
        check_history(db)
        check_multi_sheet(db, project)
        check_unsupported(db, project)

    banner("DONE")
    print("dataflow check passed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
