"""Tests for the per website sheet of a report and the retry from a report."""

from __future__ import annotations

from app.db.models import Task, TaskAttempt, TaskEvent
from tests.conftest import API

#: The website the opportunities of the tests live on.
DIRECTORY_URL = "https://report-directory.example/add-your-listing"
#: A second website so the per website sheet has more than one row.
SECOND_DIRECTORY_URL = "https://other-directory.example/submit"


def _create_opportunity(
    client, project_id: str, url: str, name: str = "Report directory"
) -> str:
    response = client.post(
        f"{API}/opportunities",
        json={
            "name": name,
            "opportunity_url": url,
            "category": "BUSINESS_DIRECTORY",
            "project_id": project_id,
        },
    )
    assert response.status_code == 201, response.text
    return response.json()["id"]


def _queue(client, opportunity_id: str) -> str:
    response = client.post(
        f"{API}/queue", json={"opportunity_ids": [opportunity_id], "type": "SUBMIT"}
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["created"] == 1, body
    return body["task_ids"][0]


def _fail_task(db_session, task_id: str, error: str = "the form was rejected") -> None:
    """Mark a task as failed - the way the worker does it."""
    task = db_session.get(Task, task_id)
    task.status = "FAILED"
    task.last_error = error
    task.attempts = (task.attempts or 0) + 1
    db_session.add(
        TaskAttempt(
            task_id=task_id,
            attempt_number=task.attempts,
            outcome="FAILED",
            error=error,
            worker="test",
        )
    )
    db_session.commit()


def _set_status(db_session, task_id: str, status: str) -> None:
    task = db_session.get(Task, task_id)
    task.status = status
    db_session.commit()


def _website_rows(workbook) -> dict[str, dict[str, object]]:
    """The rows of the "Per website" sheet, keyed by domain."""
    sheet = workbook["Per website"]
    rows = list(sheet.iter_rows(values_only=True))
    header = [str(cell) for cell in rows[0]]
    result: dict[str, dict[str, object]] = {}
    for row in rows[1:]:
        values = {header[index]: row[index] for index in range(len(header))}
        result[str(values["Domain"])] = values
    return result


def test_the_report_has_a_per_website_sheet(authenticated_client, project_id) -> None:
    """The sheet lists every website with the counters of the work."""
    import openpyxl

    opportunity_id = _create_opportunity(authenticated_client, project_id, DIRECTORY_URL)
    _queue(authenticated_client, opportunity_id)

    job = authenticated_client.post(
        f"{API}/reports", json={"project_id": project_id, "format": "xlsx"}
    ).json()
    assert job["status"] == "COMPLETED", job

    workbook = openpyxl.load_workbook(job["file_path"])
    assert "Per website" in workbook.sheetnames, workbook.sheetnames
    # The sheet comes right after the opportunities.
    assert workbook.sheetnames[:3] == [
        "Executive Summary",
        "All Opportunities",
        "Per website",
    ], workbook.sheetnames

    rows = _website_rows(workbook)
    domain = "report-directory.example"
    assert domain in rows, rows
    row = rows[domain]
    assert row["Website"] == "Report directory"
    assert row["Category"] == "BUSINESS_DIRECTORY"
    assert row["Opportunities"] == 1
    # The task is queued - it has not been submitted yet.
    assert row["Submitted"] == 0
    assert row["Failed"] == 0
    assert row["Pending"] == 1
    assert row["Verified"] == 0
    assert row["Waiting for details"] == 0
    assert row["Last attempt"] in (None, "")


def test_the_per_website_sheet_counts_every_state(authenticated_client, project_id, db_session) -> None:
    """Every task state is counted in its own column."""
    import openpyxl

    opportunity_id = _create_opportunity(authenticated_client, project_id, DIRECTORY_URL)
    task_id = _queue(authenticated_client, opportunity_id)
    _fail_task(db_session, task_id)

    second_id = _create_opportunity(
        authenticated_client, project_id, SECOND_DIRECTORY_URL, "Other directory"
    )
    second_task = _queue(authenticated_client, second_id)
    _set_status(db_session, second_task, "VERIFIED_LIVE")

    job = authenticated_client.post(
        f"{API}/reports", json={"project_id": project_id, "format": "xlsx"}
    ).json()
    assert job["status"] == "COMPLETED", job

    workbook = openpyxl.load_workbook(job["file_path"])
    rows = _website_rows(workbook)

    failed = rows["report-directory.example"]
    assert failed["Opportunities"] == 1
    assert failed["Failed"] == 1
    assert failed["Pending"] == 0
    assert failed["Submitted"] == 0

    verified = rows["other-directory.example"]
    assert verified["Opportunities"] == 1
    assert verified["Verified"] == 1
    assert verified["Submitted"] == 1
    assert verified["Failed"] == 0


def test_the_csv_export_contains_the_website_rows(authenticated_client, project_id) -> None:
    """The flat CSV carries the per website rows as well."""
    opportunity_id = _create_opportunity(authenticated_client, project_id, DIRECTORY_URL)
    _queue(authenticated_client, opportunity_id)

    job = authenticated_client.post(
        f"{API}/reports", json={"project_id": project_id, "format": "csv"}
    ).json()
    assert job["status"] == "COMPLETED", job

    content = authenticated_client.get(f"{API}/reports/{job['id']}/download").text
    assert "report-directory.example" in content
    # The header of the per website columns is part of the file.
    assert "Submitted" in content
    assert "Waiting for details" in content


def test_the_retry_requeues_the_failed_tasks_only(
    authenticated_client, project_id, db_session
) -> None:
    """Only the failed tasks are queued again - the finished ones stay untouched."""
    opportunity_id = _create_opportunity(authenticated_client, project_id, DIRECTORY_URL)
    failed_task = _queue(authenticated_client, opportunity_id)
    _fail_task(db_session, failed_task)

    second_id = _create_opportunity(
        authenticated_client, project_id, SECOND_DIRECTORY_URL, "Other directory"
    )
    verified_task = _queue(authenticated_client, second_id)
    _set_status(db_session, verified_task, "VERIFIED_LIVE")

    job = authenticated_client.post(
        f"{API}/reports", json={"project_id": project_id, "format": "xlsx"}
    ).json()
    assert job["status"] == "COMPLETED", job

    response = authenticated_client.post(f"{API}/reports/{job['id']}/retry-failed", json={})
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["retried"] == 1, body
    assert body["skipped"] == 0, body
    assert body["task_ids"] == [failed_task], body

    db_session.expire_all()
    retried = db_session.get(Task, failed_task)
    assert retried.status == "QUEUED"
    # The attempt history is not thrown away.
    attempts = db_session.query(TaskAttempt).filter_by(task_id=failed_task).count()
    assert attempts == 1

    # The verified task is untouched.
    verified = db_session.get(Task, verified_task)
    assert verified.status == "VERIFIED_LIVE"


def test_the_retry_writes_an_activity_event(authenticated_client, project_id, db_session) -> None:
    """Every re-queued task leaves a "task.retry" event behind."""
    opportunity_id = _create_opportunity(authenticated_client, project_id, DIRECTORY_URL)
    task_id = _queue(authenticated_client, opportunity_id)
    _fail_task(db_session, task_id)

    job = authenticated_client.post(
        f"{API}/reports", json={"project_id": project_id, "format": "xlsx"}
    ).json()
    authenticated_client.post(f"{API}/reports/{job['id']}/retry-failed", json={})

    db_session.expire_all()
    events = db_session.query(TaskEvent).filter_by(task_id=task_id, event_type="task.retry").count()
    assert events == 1


def test_a_task_that_was_already_retried_is_skipped(
    authenticated_client, project_id, db_session
) -> None:
    """A task that is on its way already is counted as skipped, not queued again."""
    opportunity_id = _create_opportunity(authenticated_client, project_id, DIRECTORY_URL)
    task_id = _queue(authenticated_client, opportunity_id)
    _fail_task(db_session, task_id)
    # The task was retried by hand and is running now.
    _set_status(db_session, task_id, "RUNNING")

    job = authenticated_client.post(
        f"{API}/reports", json={"project_id": project_id, "format": "xlsx"}
    ).json()

    body = authenticated_client.post(f"{API}/reports/{job['id']}/retry-failed", json={}).json()
    assert body["retried"] == 0, body
    assert body["skipped"] == 1, body
    assert body["task_ids"] == [], body

    db_session.expire_all()
    assert db_session.get(Task, task_id).status == "RUNNING"


def test_the_retry_does_not_create_a_second_task(
    authenticated_client, project_id, db_session
) -> None:
    """A second call does not queue the work a second time."""
    opportunity_id = _create_opportunity(authenticated_client, project_id, DIRECTORY_URL)
    task_id = _queue(authenticated_client, opportunity_id)
    _fail_task(db_session, task_id)

    job = authenticated_client.post(
        f"{API}/reports", json={"project_id": project_id, "format": "xlsx"}
    ).json()

    first = authenticated_client.post(
        f"{API}/reports/{job['id']}/retry-failed", json={}
    ).json()
    assert first["retried"] == 1, first

    second = authenticated_client.post(
        f"{API}/reports/{job['id']}/retry-failed", json={}
    ).json()
    assert second["retried"] == 0, second
    assert second["task_ids"] == [], second
    # The task that was already retried is reported as skipped.
    assert second["skipped"] == 1, second

    db_session.expire_all()
    tasks = db_session.query(Task).filter_by(opportunity_id=opportunity_id).all()
    assert len(tasks) == 1
    assert tasks[0].id == task_id
    assert tasks[0].status == "QUEUED"


def test_the_retry_can_be_limited_to_one_domain(authenticated_client, project_id, db_session) -> None:
    """The domain filter narrows the retry down to one website."""
    first_id = _create_opportunity(authenticated_client, project_id, DIRECTORY_URL)
    first_task = _queue(authenticated_client, first_id)
    _fail_task(db_session, first_task)

    second_id = _create_opportunity(
        authenticated_client, project_id, SECOND_DIRECTORY_URL, "Other directory"
    )
    second_task = _queue(authenticated_client, second_id)
    _fail_task(db_session, second_task)

    job = authenticated_client.post(
        f"{API}/reports", json={"project_id": project_id, "format": "xlsx"}
    ).json()

    response = authenticated_client.post(
        f"{API}/reports/{job['id']}/retry-failed",
        json={"domain": "other-directory.example"},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["retried"] == 1, body
    assert body["task_ids"] == [second_task], body

    db_session.expire_all()
    assert db_session.get(Task, first_task).status == "FAILED"
    assert db_session.get(Task, second_task).status == "QUEUED"


def test_the_retry_of_an_unknown_report_is_404(authenticated_client) -> None:
    response = authenticated_client.post(
        f"{API}/reports/does-not-exist/retry-failed", json={}
    )
    assert response.status_code == 404
