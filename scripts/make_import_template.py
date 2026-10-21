#!/usr/bin/env python3
"""Create the Excel template for the import centre.

The file contains the column names the importer recognises, one example row
and a second sheet that documents every column.

Usage:
    .venv/bin/python scripts/make_import_template.py [--out path.xlsx]
"""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = [
    ("Website", 30, "Name of the directory / site"),
    ("Opportunity URL", 46, "The exact page that accepts a listing"),
    ("Category", 22, "BUSINESS_DIRECTORY, INDUSTRY_DIRECTORY, GUEST_POST, ..."),
    ("Status", 16, "Leave empty for a new site - PENDING / FAILED / VERIFIED_LIVE are understood"),
    ("Anchor Text", 22, "The clickable text of the link"),
    ("Target URL", 40, "Your page that should receive the link"),
    ("Description", 46, "One or two sentences about the business"),
    ("Email", 30, "Contact address of the directory"),
    ("Project", 22, "Name of the project (optional)"),
    ("Notes", 30, "Anything worth remembering"),
]

EXAMPLE = [
    "Example Directory",
    "https://example.com/add-your-business",
    "BUSINESS_DIRECTORY",
    "",
    "Servchip",
    "https://servchip.com",
    "Servchip supplies AI chips, GPUs and accelerators with certified procurement and engineering support.",
    "listings@example.com",
    "Servchip backlinks",
    "Checked by phone",
]


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", default="import-template.xlsx")
    args = parser.parse_args()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Opportunities"
    header_fill = PatternFill("solid", fgColor="4F46E5")
    for index, (name, width, _hint) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = width
    for index, value in enumerate(EXAMPLE, start=1):
        sheet.cell(row=2, column=index, value=value)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"

    docs = workbook.create_sheet("Column guide")
    docs.cell(row=1, column=1, value="Column").font = Font(bold=True)
    docs.cell(row=1, column=2, value="What to put there").font = Font(bold=True)
    docs.column_dimensions["A"].width = 22
    docs.column_dimensions["B"].width = 70
    for index, (name, _width, hint) in enumerate(COLUMNS, start=2):
        docs.cell(row=index, column=1, value=name)
        docs.cell(row=index, column=2, value=hint)

    target = Path(args.out)
    workbook.save(target)
    print(f"Wrote {target.resolve()}")
    print("Fill the first sheet, then drag the file into Import Centre -> Preview -> Commit.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
